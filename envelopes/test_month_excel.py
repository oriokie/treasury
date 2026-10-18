"""Month multi-sheet Excel: one sheet per Sabbath, no Channel column."""
import datetime as dt
import io
from decimal import Decimal

import openpyxl
from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.urls import reverse

from core.roles import TREASURER
from core.utils import saturdays_of_month
from departments.models import Department
from envelopes.models import Envelope

SAB1 = dt.date(2026, 6, 6)
SAB2 = dt.date(2026, 6, 13)


def _user(username, role):
    u = User.objects.create_user(username, password="x")
    u.groups.add(Group.objects.get_or_create(name=role)[0])
    return u


class MonthExcelTests(TestCase):
    def setUp(self):
        self.tr = _user("mx_tr", TREASURER)
        self.client.force_login(self.tr)
        self.tithe = Department.objects.create(name="MxTithe", fund_type="TRUST")
        self.lcb = Department.objects.create(name="MxLCB", fund_type="LOCAL")
        self._env("MX-1", "ALPHA", SAB1, tithe=100, lcb=50)
        self._env("MX-2", "BETA", SAB2, tithe=200, channel="BANK")

    def _env(self, receipt, name, date, tithe=0, lcb=0, channel="CASH"):
        e = Envelope.objects.create(
            date=date, receipt_no=receipt, contributor_name=name,
            channel=channel, total=Decimal(tithe) + Decimal(lcb),
            recorded_by=self.tr)
        if tithe:
            e.lines.create(department=self.tithe, amount=Decimal(tithe))
        if lcb:
            e.lines.create(department=self.lcb, amount=Decimal(lcb))
        return e

    def test_list_page_links_month_excel(self):
        r = self.client.get(reverse("envelope_list") + "?month=2026-06")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, reverse("envelope_month_excel"))
        self.assertContains(r, "Excel (all Sabbaths)")

    def test_month_workbook_has_sheet_per_sabbath_without_channel(self):
        r = self.client.get(reverse("envelope_month_excel") + "?month=2026-06")
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheet", r["Content-Type"])
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        expected = [s.strftime("%d %b %Y") for s in saturdays_of_month(2026, 6)]
        self.assertEqual(wb.sheetnames, expected)

        sab1 = wb[SAB1.strftime("%d %b %Y")]
        header = next(row for row in sab1.iter_rows(min_row=1, max_row=5, values_only=True)
                      if row and row[0] == "No")
        self.assertIn("Contributor", header)
        self.assertIn("Receipt", header)
        self.assertNotIn("Channel", header)
        self.assertIn("MxTithe", header)
        # data + summary present
        values = [row[1] for row in sab1.iter_rows(values_only=True)]
        self.assertIn("ALPHA", values)
        self.assertIn("SUMMARY", values)
        self.assertIn("GRAND TOTAL GIVEN", values)

        sab2 = wb[SAB2.strftime("%d %b %Y")]
        values2 = [row[1] for row in sab2.iter_rows(values_only=True)]
        self.assertIn("BETA", values2)

    def test_single_sabbath_excel_still_includes_channel(self):
        r = self.client.get(reverse("envelope_sabbath_excel") + f"?date={SAB1.isoformat()}")
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        header = next(row for row in wb.active.iter_rows(
            min_row=1, max_row=5, values_only=True) if row and row[0] == "No")
        self.assertIn("Channel", header)
