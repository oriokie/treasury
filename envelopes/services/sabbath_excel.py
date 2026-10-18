"""Build Treasurer's Cash Statement worksheets for one or more Sabbaths."""
from __future__ import annotations

import datetime as dt
from decimal import Decimal

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from departments.models import numbered_subgroup_parent_map
from giving.models import SplitFund


def _strip_suffix(receipt):
    # receipts are stored globally-unique as "<MON><sidx>-<original>"
    # (e.g. "JUN1-0421"); show only the original number to the reader.
    s = str(receipt or "")
    if "-" in s:
        head, _, tail = s.partition("-")
        if head[:3].isalpha() and any(ch.isdigit() for ch in head):
            return tail or s
    return s


def _entry_columns(envs):
    """Collapsed display columns + real-fund list for the summary block."""
    split_parent = {}
    split_name_is_trust = {}
    for sf in SplitFund.objects.filter(active=True).prefetch_related(
            "components__department"):
        for comp in sf.components.all():
            split_parent[comp.department_id] = sf.name
            if comp.department.is_trust:
                split_name_is_trust[sf.name] = True

    subgroup_parent = numbered_subgroup_parent_map()

    present = {}
    for e in envs:
        for l in e.lines.all():
            if l.department_id in split_parent:
                name = split_parent[l.department_id]
                present.setdefault(("split", name),
                                   split_name_is_trust.get(name, False))
            else:
                dept = subgroup_parent.get(l.department_id, l.department)
                present.setdefault(("dept", dept.id), dept)

    entry_cols = []
    for key, val in present.items():
        if key[0] == "split":
            nm = key[1]
            entry_cols.append({
                "label": nm, "is_trust": val,
                "ids": {d for d, n in split_parent.items() if n == nm}})
        else:
            d = val
            ids = {d.id} | {cid for cid, p in subgroup_parent.items()
                            if p.id == d.id}
            entry_cols.append({
                "label": d.name, "is_trust": d.is_trust, "ids": ids})
    entry_cols.sort(key=lambda c: (not c["is_trust"], c["label"]))

    present_real = {}
    for e in envs:
        for l in e.lines.all():
            dept = subgroup_parent.get(l.department_id, l.department)
            present_real[dept.id] = dept
    funds = sorted(present_real.values(), key=lambda d: (not d.is_trust, d.name))
    return entry_cols, funds, subgroup_parent


def write_sabbath_sheet(ws, sab, envs, church, *, include_channel=True):
    """Fill ``ws`` with one Sabbath statement (same layout as the single download).

    When ``include_channel`` is False the Channel column is omitted — used by
    the month multi-sheet workbook.
    """
    entry_cols, funds, subgroup_parent = _entry_columns(envs)

    bold = Font(bold=True)
    white = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F5F4F")
    amber = PatternFill("solid", fgColor="B07D2C")
    grey = PatternFill("solid", fgColor="E8E2D4")
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    meta = ["No", "Contributor", "Receipt"]
    if include_channel:
        meta.append("Channel")
    ncols = len(meta) + len(entry_cols) + 1
    first_fund_col = len(meta) + 1  # 1-based

    ws.append([church])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ncols, 1))
    ws.cell(1, 1).font = Font(bold=True, size=15, color="1F5F4F")
    ws.cell(1, 1).alignment = Alignment(horizontal="center")
    ws.append([f"Treasurer's Cash Statement — Sabbath {sab:%d %B %Y}"])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(ncols, 1))
    ws.cell(2, 1).font = Font(bold=True, size=12)
    ws.cell(2, 1).alignment = Alignment(horizontal="center")

    header = meta + [c["label"] for c in entry_cols] + ["Total"]
    ws.append(header)
    hr = ws.max_row
    for c in range(1, len(header) + 1):
        cell = ws.cell(hr, c)
        cell.font = white
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = border

    for i, e in enumerate(envs, start=1):
        amt = {}
        for l in e.lines.all():
            amt[l.department_id] = amt.get(l.department_id, Decimal(0)) + l.amount
        row_cells = []
        for col in entry_cols:
            v = sum((amt.get(did, Decimal(0)) for did in col["ids"]), Decimal(0))
            row_cells.append(float(v) if v else "")
        front = [i, e.contributor_name, _strip_suffix(e.receipt_no)]
        if include_channel:
            front.append(e.get_channel_display())
        ws.append(front + row_cells + [float(e.total)])

    totals = {f.id: Decimal(0) for f in funds}
    grand = Decimal(0)
    for e in envs:
        for l in e.lines.all():
            did = subgroup_parent.get(l.department_id, l.department).id
            if did in totals:
                totals[did] += l.amount
        grand += e.total
    col_totals = []
    for col in entry_cols:
        col_totals.append(float(sum((totals.get(did, Decimal(0))
                                     for did in col["ids"]), Decimal(0))))
    total_front = ["", "TOTAL", ""]
    if include_channel:
        total_front.append("")
    ws.append(total_front + col_totals + [float(grand)])
    last_data = ws.max_row
    for c in range(1, len(header) + 1):
        ws.cell(last_data, c).font = bold
        ws.cell(last_data, c).fill = grey

    for r in range(hr, last_data + 1):
        for c in range(1, len(header) + 1):
            cell = ws.cell(r, c)
            cell.border = border
            if c >= first_fund_col and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"

    ws.append([])
    summary_start = ws.max_row + 1
    ws.append(["", "SUMMARY"])
    ws.cell(ws.max_row, 2).font = Font(bold=True, size=12)
    ws.append(["", "Trust fund items (remitted)"])
    ws.cell(ws.max_row, 2).font = bold
    ws.cell(ws.max_row, 2).fill = amber
    ws.cell(ws.max_row, 3).fill = amber
    trust_total = Decimal(0)
    for f in funds:
        if f.is_trust:
            ws.append(["", f.name, float(totals[f.id])])
            trust_total += totals[f.id]
    ws.append(["", "Total trust funds", float(trust_total)])
    ws.cell(ws.max_row, 2).font = bold
    ws.cell(ws.max_row, 3).font = bold
    ws.cell(ws.max_row, 2).fill = grey
    ws.cell(ws.max_row, 3).fill = grey
    ws.append(["", "Local fund items (retained)"])
    ws.cell(ws.max_row, 2).font = bold
    ws.cell(ws.max_row, 2).fill = amber
    ws.cell(ws.max_row, 3).fill = amber
    local_total = Decimal(0)
    for f in funds:
        if not f.is_trust:
            ws.append(["", f.name, float(totals[f.id])])
            local_total += totals[f.id]
    ws.append(["", "Total local funds", float(local_total)])
    ws.cell(ws.max_row, 2).font = bold
    ws.cell(ws.max_row, 3).font = bold
    ws.cell(ws.max_row, 2).fill = grey
    ws.cell(ws.max_row, 3).fill = grey
    ws.append(["", "GRAND TOTAL GIVEN", float(grand)])
    ws.cell(ws.max_row, 2).font = Font(bold=True, size=12)
    ws.cell(ws.max_row, 3).font = Font(bold=True, size=12)
    summary_end = ws.max_row

    for r in range(summary_start, summary_end + 1):
        for c in (2, 3):
            cell = ws.cell(r, c)
            cell.border = border
            if c == 3 and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 14
    if include_channel:
        ws.column_dimensions["D"].width = 12
    for i in range(len(meta), len(header)):
        col = get_column_letter(i + 1)
        ws.column_dimensions[col].width = max(
            ws.column_dimensions[col].width or 10, 13)
    ws.freeze_panes = "A4"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.6, bottom=0.5,
                                  header=0.3, footer=0.3)
    ws.print_title_rows = "1:3"
    ws.oddFooter.center.text = f"{church} — Sabbath {sab:%d %b %Y} — Page &P of &N"


def sheet_title_for(sab: dt.date) -> str:
    """Excel sheet name (≤31 chars) for a Sabbath date."""
    return sab.strftime("%d %b %Y")
