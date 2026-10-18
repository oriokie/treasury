"""Build and send fund-leader period statements by SMS and email.

Figures come from ``reports.services.balances.department_summary`` — the same
opening / receipts / expenses / closing the fund reports use.

Rules for a leader linked to a parent and its children:
  * **SMS** — one message with the *consolidated* total (parent + sub-accounts).
  * **Email** — one message with a single Excel attachment: a summary sheet
    (main + sub-accounts, like ``/reports/fund/<id>/?export=subgroups``) plus
    a detailed ledger sheet of entries for each account (same rows as the fund
    ledger export).

Recipients without a valid phone (SMS) or email (email) are skipped.
"""
from __future__ import annotations

import datetime as dt
import io
from decimal import Decimal

from django.core.exceptions import ValidationError
from django.core.validators import validate_email

from members.models import normalize_phone
from reports.services import balances


def _money(amount):
    return f"{(amount or Decimal(0)):,.0f}"


def _period_label(start, end):
    if start.year == end.year and start.month == end.month:
        return f"{start:%d}–{end:%d %b %Y}"
    return f"{start:%d %b %Y} – {end:%d %b %Y}"


def _zero_figures(dept=None):
    return {
        "opening": Decimal(0), "receipts": Decimal(0),
        "expenses": Decimal(0), "transfers_in": Decimal(0),
        "transfers_out": Decimal(0), "closing": Decimal(0),
        "department": dept,
    }


def statement_rows_by_id(start, end):
    """``{dept_id: row}`` for every fund that has something to report."""
    return {
        r["department"].id: r
        for r in balances.department_summary(start, end, consolidated=False)
    }


def figures_for(dept, rows_by_id):
    """Opening / receipts / expenses / closing for one fund, or zeros."""
    r = rows_by_id.get(dept.id)
    if not r:
        out = _zero_figures(dept)
        return out
    return r


def report_roots(leader):
    """Linked funds to report, collapsing children when the parent is linked.

    If a leader is assigned both a parent and its sub-accounts, only the parent
    appears here — sub-accounts roll into its consolidated figures / Excel.
    """
    funds = list(leader.departments.filter(active=True)
                 .select_related("parent")
                 .prefetch_related("subgroups")
                 .order_by("name"))
    linked = {f.id for f in funds}
    return [f for f in funds if f.parent_id not in linked]


def consolidate_figures(root, rows_by_id):
    """Parent + all active sub-accounts rolled into one O/R/E/C set."""
    parts = [root] + list(root.subgroups.filter(active=True).order_by("name"))
    out = _zero_figures(root)
    for dept in parts:
        fig = figures_for(dept, rows_by_id)
        for key in ("opening", "receipts", "expenses",
                    "transfers_in", "transfers_out"):
            out[key] += fig.get(key) or Decimal(0)
    out["closing"] = (out["opening"] + out["receipts"] - out["expenses"]
                      + out["transfers_in"] - out["transfers_out"])
    return out


def sum_figures(figs_list, dept=None):
    out = _zero_figures(dept)
    for fig in figs_list:
        for key in ("opening", "receipts", "expenses",
                    "transfers_in", "transfers_out", "closing"):
            out[key] += fig.get(key) or Decimal(0)
    return out


def sms_body(leader, fund_label, figures, start, end, church=""):
    from core.services.sms import sms_first_name
    who = sms_first_name(leader.name) or leader.name
    lines = [
        f"{who}, {fund_label} ({_period_label(start, end)})",
        f"Opening: {_money(figures['opening'])}",
        f"Receipts: {_money(figures['receipts'])}",
        f"Expenses: {_money(figures['expenses'])}",
        f"Closing: {_money(figures['closing'])}",
    ]
    if church:
        lines.append(f"- {church}")
    return "\n".join(lines)


def email_body(leader, fund_labels, figures, start, end, church=""):
    net_t = (figures.get("transfers_in") or Decimal(0)) - (
        figures.get("transfers_out") or Decimal(0))
    label = ", ".join(fund_labels) if fund_labels else "your funds"
    lines = [
        f"Dear {leader.name},",
        "",
        f"Fund statement — {label}",
        f"Period: {_period_label(start, end)}",
        "",
        "Consolidated totals (main fund + sub-accounts):",
        f"Opening balance:  {_money(figures['opening'])}",
        f"Receipts:         {_money(figures['receipts'])}",
        f"Expenses:         {_money(figures['expenses'])}",
    ]
    if net_t:
        lines.append(f"Net transfers:    {_money(net_t)}")
    lines += [
        f"Closing balance:  {_money(figures['closing'])}",
        "",
        "A spreadsheet is attached: a summary of the main fund and each "
        "sub-account, plus a detailed statement sheet of entries for every account.",
        "",
    ]
    if church:
        lines.append(f"{church}")
        lines.append("Church treasury")
    return "\n".join(lines)


def email_subject(fund_labels, start, end, church=""):
    prefix = f"{church} — " if church else ""
    if len(fund_labels) == 1:
        name = fund_labels[0]
    elif fund_labels:
        name = f"{len(fund_labels)} funds"
    else:
        name = "Fund"
    return f"{prefix}{name} statement ({_period_label(start, end)})"


def _safe_sheet_title(name):
    cleaned = "".join(c if c not in ":\\/?*[]" else "-" for c in (name or "Fund"))
    return (cleaned or "Fund")[:31]


def _unique_sheet_title(used, name):
    title = _safe_sheet_title(name)
    base, n = title, 2
    while title in used:
        suffix = f"-{n}"
        title = _safe_sheet_title(base[: 31 - len(suffix)] + suffix)
        n += 1
    used.add(title)
    return title


def ledger_entries_for(dept, start, end):
    """Detailed fund ledger rows for one account — same basis as FundLedgerView.

    Returns ``(opening, entries, closing)`` where each entry has date, src,
    desc, debit, credit, balance.
    """
    from cashbook.models import Expense, ExpenseRefund, FundTransfer
    from giving.models import Transaction
    from reports.services.balances import brought_forward

    receipts = (Transaction.objects.confirmed_credits()
                .filter(department=dept, date__gte=start, date__lte=end)
                .order_by("date"))
    payments = (Expense.objects.filter(
        department=dept, date__gte=start, date__lte=end,
        status__in=[Expense.Status.APPROVED, Expense.Status.PAID])
        .order_by("date"))
    entries = []
    for t in receipts:
        if t.excluded_from_income:
            base = t.payer_name or t.reference or "Receipt"
            entries.append({
                "date": t.date,
                "desc": f"{base} — loan / financing receipt (not income)",
                "credit": t.amount, "debit": None, "src": "Financing",
            })
        else:
            entries.append({
                "date": t.date,
                "desc": t.payer_name or t.reference or "Receipt",
                "credit": t.amount, "debit": None, "src": "Receipt",
            })
    for x in payments:
        entries.append({
            "date": x.date, "desc": x.description,
            "credit": None, "debit": x.amount, "src": "Expense",
        })
    for rf in (ExpenseRefund.objects.filter(
            expense__department=dept, date__gte=start, date__lte=end,
            expense__status__in=[Expense.Status.APPROVED, Expense.Status.PAID])
            .select_related("expense")):
        entries.append({
            "date": rf.date,
            "desc": f"Refund — {rf.expense.description}",
            "credit": rf.amount, "debit": None, "src": "Refund",
        })
    for tr in FundTransfer.objects.filter(
            destination=dept, date__gte=start, date__lte=end):
        entries.append({
            "date": tr.date,
            "desc": f"Transfer from {tr.source.name}"
                    + (f" — {tr.reason}" if tr.reason else ""),
            "credit": tr.amount, "debit": None, "src": "Transfer",
        })
    for tr in FundTransfer.objects.filter(
            source=dept, date__gte=start, date__lte=end):
        entries.append({
            "date": tr.date,
            "desc": f"Transfer to {tr.destination.name}"
                    + (f" — {tr.reason}" if tr.reason else ""),
            "credit": None, "debit": tr.amount, "src": "Transfer",
        })
    entries.sort(key=lambda r: r["date"])
    opening = brought_forward(dept, start)
    running = opening
    for en in entries:
        running += (en["credit"] or 0) - (en["debit"] or 0)
        en["balance"] = running
    return opening, entries, running


def _write_header_block(ws, church, title_line, styles):
    """Church + title rows; returns the next free row index (1-based)."""
    Font = styles["Font"]
    r = 1
    if church:
        ws.cell(row=r, column=1, value=church).font = Font(bold=True, size=14)
        r += 1
    ws.cell(row=r, column=1, value=title_line).font = Font(
        bold=True, size=12, color="1F5F4F")
    return r + 2


def _style_header_row(ws, row, headers, styles):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = styles["white"]
        cell.fill = styles["head_fill"]
        cell.alignment = styles["Alignment"](horizontal="center")


def _write_summary_sheet(wb, root, rows_by_id, start, end, church, used, styles):
    from openpyxl.utils import get_column_letter

    title = _unique_sheet_title(used, f"{root.name} · Summary")
    ws = wb.create_sheet(title=title)
    header = ["Fund / sub-account", "Type", "Opening", "Receipts",
              "Payments", "Closing balance"]
    r = _write_header_block(
        ws, church,
        (f"{root.name} — main fund & sub-accounts "
         f"({start:%d %b %Y} to {end:%d %b %Y})"),
        styles)
    head_row = r
    _style_header_row(ws, head_row, header, styles)

    parent_fig = figures_for(root, rows_by_id)
    accounts = [(f"{root.name} (main)", root, parent_fig)]
    for sub in root.subgroups.filter(active=True).order_by("name"):
        accounts.append((sub.name, sub, figures_for(sub, rows_by_id)))

    cons = consolidate_figures(root, rows_by_id)
    rr = head_row
    for label, dept, fig in accounts:
        rr += 1
        closing = fig["closing"] if fig.get("closing") is not None else (
            fig["opening"] + fig["receipts"] - fig["expenses"]
            + fig.get("transfers_in", 0) - fig.get("transfers_out", 0))
        vals = [
            label,
            "Trust" if dept.is_trust else "Local",
            float(fig["opening"] or 0),
            float(fig["receipts"] or 0),
            float(fig["expenses"] or 0),
            float(closing),
        ]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=rr, column=c, value=v)
            if isinstance(v, float):
                cell.number_format = "#,##0.00"

    rr += 1
    for c, v in enumerate([
        "TOTAL (consolidated)", "",
        float(cons["opening"]), float(cons["receipts"]),
        float(cons["expenses"]), float(cons["closing"]),
    ], start=1):
        cell = ws.cell(row=rr, column=c, value=v)
        cell.font = styles["bold"]
        cell.fill = styles["grey"]
        if isinstance(v, float):
            cell.number_format = "#,##0.00"

    for c, h in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(c)].width = max(14, len(str(h)) + 2)
    ws.column_dimensions["A"].width = 28
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    return accounts


def _write_detail_sheet(wb, dept, label, start, end, church, used, styles):
    """One sheet of ledger entries for a single account."""
    from openpyxl.utils import get_column_letter

    title = _unique_sheet_title(used, label)
    ws = wb.create_sheet(title=title)
    header = ["Type", "Date", "Description", "Debit", "Credit", "Balance"]
    r = _write_header_block(
        ws, church,
        (f"{dept.name} — detailed statement "
         f"({start:%d %b %Y} to {end:%d %b %Y})"),
        styles)
    head_row = r
    _style_header_row(ws, head_row, header, styles)

    opening, entries, closing = ledger_entries_for(dept, start, end)
    rr = head_row + 1
    for c, v in enumerate(
            ["", "", "Opening balance", "", "", float(opening)], start=1):
        cell = ws.cell(row=rr, column=c, value=v)
        cell.font = styles["bold"]
        if isinstance(v, float):
            cell.number_format = "#,##0.00"

    for en in entries:
        rr += 1
        vals = [
            en["src"],
            en["date"].isoformat() if en["date"] else "",
            en["desc"],
            float(en["debit"]) if en["debit"] else "",
            float(en["credit"]) if en["credit"] else "",
            float(en["balance"]),
        ]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=rr, column=c, value=v)
            if isinstance(v, float):
                cell.number_format = "#,##0.00"

    rr += 1
    for c, v in enumerate(
            ["", "", "Closing balance", "", "", float(closing)], start=1):
        cell = ws.cell(row=rr, column=c, value=v)
        cell.font = styles["bold"]
        cell.fill = styles["grey"]
        if isinstance(v, float):
            cell.number_format = "#,##0.00"

    widths = {"A": 12, "B": 12, "C": 40, "D": 12, "E": 12, "F": 14}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for c, h in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(c)].width = max(
            ws.column_dimensions[get_column_letter(c)].width or 10,
            len(str(h)) + 2)
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)


def build_statement_xlsx(leader, roots, rows_by_id, start, end, church=""):
    """Excel bytes: summary sheet per root, then a detail sheet per account."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    styles = {
        "Font": Font,
        "Alignment": Alignment,
        "bold": Font(bold=True),
        "white": Font(bold=True, color="FFFFFF"),
        "head_fill": PatternFill("solid", fgColor="1F5F4F"),
        "grey": PatternFill("solid", fgColor="E8E2D4"),
    }
    used_titles = set()

    for root in roots:
        accounts = _write_summary_sheet(
            wb, root, rows_by_id, start, end, church, used_titles, styles)
        # Detailed statement sheet for the main fund and each sub-account.
        for label, dept, _fig in accounts:
            detail_label = (f"{dept.name} · Detail" if dept.id == root.id
                            else f"{dept.name}")
            _write_detail_sheet(
                wb, dept, detail_label, start, end, church, used_titles, styles)

    if not wb.sheetnames:
        ws = wb.create_sheet(title="Statement")
        ws.append(["No linked funds"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def send_sms_to_leader(leader, start, end, *, cfg=None, rows_by_id=None):
    """One SMS with consolidated totals across the leader's report roots.

    Returns ``(sent, skipped, failed)`` — ``sent`` is 0 or 1.
    """
    from core.models import SiteConfig
    from core.services.sms import send_sms

    cfg = cfg or SiteConfig.get()
    phone = normalize_phone(leader.phone)
    if not phone:
        return 0, 1, 0
    if not cfg.sms_enabled:
        return 0, 0, 1

    rows_by_id = rows_by_id or statement_rows_by_id(start, end)
    church = cfg.church_name or ""
    roots = report_roots(leader)
    if not roots:
        return 0, 1, 0

    cons_list = [consolidate_figures(r, rows_by_id) for r in roots]
    figures = sum_figures(cons_list, roots[0] if len(roots) == 1 else None)
    if len(roots) == 1:
        label = f"{roots[0].name} (incl. sub-accounts)"
    else:
        label = f"{len(roots)} funds (consolidated)"
    msg = sms_body(leader, label, figures, start, end, church)
    log = send_sms(phone, msg, cfg)
    if getattr(log, "status", "") == "SENT":
        return 1, 0, 0
    return 0, 0, 1


def send_email_to_leader(leader, start, end, *, cfg=None, rows_by_id=None):
    """One email with consolidated summary + Excel of main + sub-accounts.

    Returns ``(sent, skipped, failed)`` — ``sent`` is 0 or 1.
    """
    from core.models import SiteConfig
    from core.services.email import send_email, is_configured

    cfg = cfg or SiteConfig.get()
    email = (leader.email or "").strip()
    try:
        validate_email(email)
    except ValidationError:
        return 0, 1, 0
    if not is_configured(cfg):
        return 0, 0, 1

    rows_by_id = rows_by_id or statement_rows_by_id(start, end)
    church = cfg.church_name or ""
    roots = report_roots(leader)
    if not roots:
        return 0, 1, 0

    cons_list = [consolidate_figures(r, rows_by_id) for r in roots]
    figures = sum_figures(cons_list, roots[0] if len(roots) == 1 else None)
    labels = [r.name for r in roots]
    xlsx = build_statement_xlsx(
        leader, roots, rows_by_id, start, end, church=church)
    fname = f"fund-statement-{start:%Y%m%d}-{end:%Y%m%d}.xlsx"
    ok, _ = send_email(
        email_subject(labels, start, end, church),
        email_body(leader, labels, figures, start, end, church),
        email, cfg=cfg,
        attachments=[(
            fname, xlsx,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )])
    if ok:
        return 1, 0, 0
    return 0, 0, 1


def default_period(today=None):
    """Year-to-date ending today — matches how treasurers review fund ledgers."""
    today = today or dt.date.today()
    return dt.date(today.year, 1, 1), today
