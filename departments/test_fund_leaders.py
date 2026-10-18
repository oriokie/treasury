"""Fund Leaders directory — contacts, group/fund links, SMS and email sends."""
import datetime as dt
import io
from decimal import Decimal
from unittest.mock import patch

from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.urls import reverse

from core.roles import TREASURER
from departments.models import Department, DevelopmentGroup, FundLeader
from departments.services import fund_leaders as svc
from giving.models import Transaction


def _treasurer():
    u = User.objects.create_user("fl_tr", password="x")
    u.groups.add(Group.objects.get_or_create(name=TREASURER)[0])
    return u


class FundLeaderModelTests(TestCase):
    def test_normalises_phone_and_email_on_save(self):
        l = FundLeader.objects.create(
            name="  Jane  Leader ", phone="0712345678",
            email="Jane@Example.COM")
        self.assertEqual(l.name, "Jane Leader")
        self.assertEqual(l.phone, "254712345678")
        self.assertEqual(l.email, "jane@example.com")
        self.assertTrue(l.has_valid_phone)
        self.assertTrue(l.has_valid_email)

    def test_invalid_phone_is_flagged(self):
        l = FundLeader.objects.create(name="X", phone="123")
        self.assertFalse(l.has_valid_phone)


class FundLeaderCrudTests(TestCase):
    def setUp(self):
        self.tr = _treasurer()
        self.client.force_login(self.tr)
        self.fund = Department.objects.create(
            name="FlFund", fund_type="LOCAL")
        self.group = DevelopmentGroup.objects.create(number=7, name="G7")

    def test_list_is_in_funds_setup_nav(self):
        r = self.client.get("/")
        self.assertContains(r, reverse("fund_leader_list"))
        self.assertContains(r, "Leaders")

    def test_create_with_fund_and_group(self):
        r = self.client.post(reverse("fund_leader_create"), {
            "name": "Alice",
            "phone": "0711111111",
            "email": "alice@example.com",
            "active": "on",
            "notes": "",
            "departments": [self.fund.id],
            "groups": [self.group.id],
        })
        self.assertEqual(r.status_code, 302)
        l = FundLeader.objects.get(name="Alice")
        self.assertEqual(list(l.departments.all()), [self.fund])
        self.assertEqual(list(l.groups.all()), [self.group])

    def test_requires_phone_or_email(self):
        r = self.client.post(reverse("fund_leader_create"), {
            "name": "NoContact", "phone": "", "email": "", "active": "on",
        })
        self.assertEqual(r.status_code, 200)
        self.assertFalse(FundLeader.objects.filter(name="NoContact").exists())


class FundLeaderSendTests(TestCase):
    def setUp(self):
        self.tr = _treasurer()
        self.client.force_login(self.tr)
        self.fund = Department.objects.create(
            name="FlSendFund", fund_type="LOCAL", opening_balance=Decimal("1000"))
        self.sub = Department.objects.create(
            name="FlSub", fund_type="LOCAL", parent=self.fund,
            opening_balance=Decimal("200"))
        Transaction.objects.create(
            date=dt.date(2026, 6, 5), amount=Decimal("500"),
            direction="CREDIT", confirmed=True, channel="CASH",
            department=self.fund, allocation_status="MANUAL")
        Transaction.objects.create(
            date=dt.date(2026, 6, 6), amount=Decimal("100"),
            direction="CREDIT", confirmed=True, channel="CASH",
            department=self.sub, allocation_status="MANUAL")
        self.leader = FundLeader.objects.create(
            name="Bob Leader", phone="0722222222",
            email="bob@example.com")
        # Parent + child both linked — must still be one SMS / one email.
        self.leader.departments.add(self.fund, self.sub)

    def test_default_period_is_year_to_date(self):
        start, end = svc.default_period(dt.date(2026, 10, 18))
        self.assertEqual(start, dt.date(2026, 1, 1))
        self.assertEqual(end, dt.date(2026, 10, 18))

    def test_report_roots_collapses_linked_children(self):
        roots = svc.report_roots(self.leader)
        self.assertEqual([r.id for r in roots], [self.fund.id])

    def test_sms_body_has_opening_receipts_expenses_closing(self):
        rows = svc.statement_rows_by_id(dt.date(2026, 6, 1), dt.date(2026, 6, 30))
        fig = svc.consolidate_figures(self.fund, rows)
        body = svc.sms_body(self.leader, "FlSendFund (incl. sub-accounts)", fig,
                            dt.date(2026, 6, 1), dt.date(2026, 6, 30), "KWS")
        self.assertIn("Opening:", body)
        self.assertIn("Receipts:", body)
        self.assertIn("Expenses:", body)
        self.assertIn("Closing:", body)
        self.assertIn("FlSendFund", body)

    @patch("core.services.sms.send_sms")
    def test_sms_skips_leaders_without_valid_phone(self, mock_sms):
        bare = FundLeader.objects.create(name="NoPhone", email="x@y.com")
        bare.departments.add(self.fund)
        sent, skipped, failed = svc.send_sms_to_leader(
            bare, dt.date(2026, 6, 1), dt.date(2026, 6, 30))
        self.assertEqual((sent, skipped, failed), (0, 1, 0))
        mock_sms.assert_not_called()

    @patch("core.services.sms.send_sms")
    def test_sms_sends_one_consolidated_message(self, mock_sms):
        from core.models import SiteConfig
        cfg = SiteConfig.get()
        cfg.sms_enabled = True
        cfg.save()

        class Log:
            status = "SENT"
        mock_sms.return_value = Log()
        sent, skipped, failed = svc.send_sms_to_leader(
            self.leader, dt.date(2026, 1, 1), dt.date(2026, 6, 30))
        self.assertEqual((sent, skipped, failed), (1, 0, 0))
        self.assertEqual(mock_sms.call_count, 1)
        body = mock_sms.call_args.args[1]
        # Consolidated receipts = 500 (parent) + 100 (sub)
        self.assertIn("Receipts: 600", body)

    @patch("core.services.email.send_email", return_value=(True, "ok"))
    @patch("core.services.email.is_configured", return_value=True)
    def test_email_sends_one_message_with_xlsx_attachment(self, _cfg, mock_email):
        sent, skipped, failed = svc.send_email_to_leader(
            self.leader, dt.date(2026, 1, 1), dt.date(2026, 6, 30))
        self.assertEqual((sent, skipped, failed), (1, 0, 0))
        self.assertEqual(mock_email.call_count, 1)
        kwargs = mock_email.call_args.kwargs
        atts = kwargs.get("attachments") or []
        self.assertEqual(len(atts), 1)
        fname, content, mime = atts[0]
        self.assertTrue(fname.endswith(".xlsx"))
        self.assertIn("spreadsheet", mime)
        self.assertGreater(len(content), 100)
        # Summary + a detailed ledger sheet per account (main and sub).
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content))
        self.assertTrue(any("Summary" in n for n in wb.sheetnames))
        summary = next(n for n in wb.sheetnames if "Summary" in n)
        values = [cell for row in wb[summary].iter_rows(values_only=True)
                  for cell in row]
        self.assertTrue(any(v and "FlSub" in str(v) for v in values))
        self.assertTrue(any(v and "consolidated" in str(v).lower() for v in values))
        detail_main = next(n for n in wb.sheetnames if "Detail" in n)
        detail_vals = [cell for row in wb[detail_main].iter_rows(values_only=True)
                       for cell in row]
        self.assertTrue(any(v == "Opening balance" for v in detail_vals))
        self.assertTrue(any(v == "Closing balance" for v in detail_vals))
        self.assertTrue(any(n == "FlSub" or n.startswith("FlSub")
                            for n in wb.sheetnames))

    @patch("core.services.sms.send_sms")
    def test_send_one_view_posts_sms(self, mock_sms):
        from core.models import SiteConfig
        cfg = SiteConfig.get()
        cfg.sms_enabled = True
        cfg.sms_api_key = "k"
        cfg.sms_partner_id = "p"
        cfg.sms_shortcode = "s"
        cfg.save()

        class Log:
            status = "SENT"
        mock_sms.return_value = Log()
        r = self.client.post(reverse("fund_leader_send_one", args=[self.leader.pk]), {
            "channel": "sms",
            "start": "2026-01-01",
            "end": "2026-06-30",
        })
        self.assertEqual(r.status_code, 302)
        self.assertEqual(mock_sms.call_count, 1)

    def test_send_page_lists_only_eligible(self):
        FundLeader.objects.create(name="NoFund", phone="0733333333")
        r = self.client.get(reverse("fund_leader_send") + "?channel=sms")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Bob Leader")
        self.assertNotContains(r, "NoFund")
        self.assertContains(r, "one SMS")
